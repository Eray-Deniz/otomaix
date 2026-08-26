"""Document processing service for RAG pipeline.

Supports PDF, Word (.docx), and Excel (.xlsx) files.
- Small docs (< 8000 tokens ≈ 32 000 chars): store full text in brand_documents.raw_text
- Large docs (>= 8000 tokens): split into 512-token chunks (≈ 2048 chars) with 50-token overlap
  and store in brand_document_chunks.

If OPENAI_API_KEY is configured, embeddings are generated for each chunk
(text-embedding-3-small, 1536 dims) to enable vector similarity search.
Otherwise chunks are stored without embeddings (fallback: first-N retrieval).
"""

import io
import re
from uuid import UUID

import asyncpg
from fastapi import HTTPException

from app.core.config import settings
from app.services.storage import r2

# Token approximation: ~4 chars per token
CHARS_PER_TOKEN = 4
SMALL_DOC_THRESHOLD = 8_000 * CHARS_PER_TOKEN  # 32 000 chars
CHUNK_SIZE_CHARS = 512 * CHARS_PER_TOKEN        # 2 048 chars
CHUNK_OVERLAP_CHARS = 50 * CHARS_PER_TOKEN      # 200 chars

ALLOWED_MIME_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
    "text/plain",
}


# ─── Text extraction ────────────────────────────────────────────────────────

def _extract_pdf(content: bytes) -> str:
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(p.strip() for p in pages if p.strip())
    except Exception:
        return ""


def _extract_docx(content: bytes) -> str:
    try:
        import docx
        doc = docx.Document(io.BytesIO(content))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)
    except Exception:
        return ""


def _extract_xlsx(content: bytes) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in wb.worksheets:
            lines.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                line = "\t".join(cells).strip()
                if line:
                    lines.append(line)
        return "\n".join(lines)
    except Exception:
        return ""


def extract_text(content: bytes, mime_type: str) -> str:
    """Extract text from document bytes based on MIME type."""
    mime = mime_type.lower()
    if mime == "application/pdf":
        return _extract_pdf(content)
    if mime in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    ):
        return _extract_docx(content)
    if mime in (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    ):
        return _extract_xlsx(content)
    if mime.startswith("text/"):
        return content.decode("utf-8", errors="replace")
    return ""


# ─── Chunking ───────────────────────────────────────────────────────────────

def _split_into_chunks(text: str) -> list[str]:
    """Split text into overlapping chunks."""
    chunks: list[str] = []
    start = 0
    length = len(text)
    while start < length:
        end = min(start + CHUNK_SIZE_CHARS, length)
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= length:
            break
        start = end - CHUNK_OVERLAP_CHARS
    return chunks


# ─── Embeddings (optional) ──────────────────────────────────────────────────

async def _generate_embedding(text: str) -> list[float] | None:
    """Generate a 1536-dim embedding using OpenAI text-embedding-3-small.
    Returns None if OPENAI_API_KEY is not configured.
    """
    if not getattr(settings, "OPENAI_API_KEY", ""):
        return None
    try:
        import openai
        client = openai.AsyncOpenAI(api_key=settings.OPENAI_API_KEY)
        resp = await client.embeddings.create(
            model="text-embedding-3-small",
            input=text,
        )
        return resp.data[0].embedding
    except Exception:
        return None


# ─── Main pipeline ──────────────────────────────────────────────────────────

async def process_document(
    document_id: UUID,
    brand_id: UUID | None,
    content: bytes,
    mime_type: str,
    db: asyncpg.Connection,
    *,
    kind: str = "brand",
) -> dict:
    """Extract text, classify size, store raw_text or chunks in DB.

    kind="brand"  → social.brand_documents + brand_document_chunks (denorm brand_id).
    kind="product" → social.product_documents + product_document_chunks (no brand_id col).

    Returns {"mode": "raw"|"chunks", "char_count": int, "chunk_count": int}
    """
    text = extract_text(content, mime_type)
    if not text:
        return {"mode": "none", "char_count": 0, "chunk_count": 0}

    char_count = len(text)

    if kind == "product":
        docs_table = "social.product_documents"
        chunks_table = "social.product_document_chunks"
    else:
        docs_table = "social.brand_documents"
        chunks_table = "social.brand_document_chunks"

    if char_count < SMALL_DOC_THRESHOLD:
        await db.execute(
            f"UPDATE {docs_table} SET raw_text = $2 WHERE id = $1",
            document_id,
            text,
        )
        return {"mode": "raw", "char_count": char_count, "chunk_count": 0}

    chunks = _split_into_chunks(text)
    for idx, chunk_text in enumerate(chunks):
        embedding = await _generate_embedding(chunk_text)
        if kind == "product":
            if embedding is not None:
                await db.execute(
                    f"""
                    INSERT INTO {chunks_table}
                        (document_id, chunk_index, content, embedding)
                    VALUES ($1, $2, $3, $4::vector)
                    """,
                    document_id,
                    idx,
                    chunk_text,
                    str(embedding),
                )
            else:
                await db.execute(
                    f"""
                    INSERT INTO {chunks_table}
                        (document_id, chunk_index, content)
                    VALUES ($1, $2, $3)
                    """,
                    document_id,
                    idx,
                    chunk_text,
                )
        else:
            if embedding is not None:
                await db.execute(
                    f"""
                    INSERT INTO {chunks_table}
                        (document_id, brand_id, chunk_index, content, embedding)
                    VALUES ($1, $2, $3, $4, $5::vector)
                    """,
                    document_id,
                    brand_id,
                    idx,
                    chunk_text,
                    str(embedding),
                )
            else:
                await db.execute(
                    f"""
                    INSERT INTO {chunks_table}
                        (document_id, brand_id, chunk_index, content)
                    VALUES ($1, $2, $3, $4)
                    """,
                    document_id,
                    brand_id,
                    idx,
                    chunk_text,
                )

    return {"mode": "chunks", "char_count": char_count, "chunk_count": len(chunks)}


# ─── RAG retrieval ──────────────────────────────────────────────────────────

async def get_document_context(
    document_ids: list[UUID],
    prompt: str,
    db: asyncpg.Connection,
    max_chars: int = 6_000,
    *,
    kind: str = "brand",
    brand_id: UUID,
) -> str:
    """Build document context string for injection into AI prompt.

    kind="brand"  → reads brand_documents + brand_document_chunks.
    kind="product" → reads product_documents + product_document_chunks (name col is `filename`).

    For small docs: include raw_text directly.
    For large docs (chunks): vector search if embeddings exist, otherwise first-N chunks.
    Returns combined context string (truncated to max_chars).

    `brand_id` ZORUNLU ve keyword-only: çağıranın DOĞRULADIĞI marka. Kapsam burada
    uygulanır, çağıranda değil — çağıranlar `payload.brand_id`'nin sahipliğini
    doğruluyordu ama `document_ids`'i doğrulamadan geçiriyordu, dolayısıyla başka
    kiracının doküman kimliğini bilen biri o dokümanın metnini kendi üretim bağlamına
    enjekte ettirebiliyordu (security review 2026-08-26, S2). Parametrenin zorunlu
    olması bilinçli: yeni bir çağıran kapsamı unutursa sessizce sızdırmaz, TypeError verir.

    İstenen kimliklerden biri bile bu markaya ait değilse İSTEĞİN TAMAMI 404 ile
    reddedilir — kısmi sonuç dönmek, saldırgana hangi kimliğin var olduğunu söyler.
    """
    if not document_ids:
        return ""

    if kind == "product":
        chunks_table = "social.product_document_chunks"
        docs_sql = """
            SELECT d.id, d.filename AS name, d.raw_text
            FROM social.product_documents d
            JOIN social.brand_products p ON p.id = d.product_id
            WHERE d.id = ANY($1::uuid[]) AND p.brand_id = $2
        """
    else:
        chunks_table = "social.brand_document_chunks"
        docs_sql = """
            SELECT id, name, raw_text
            FROM social.brand_documents
            WHERE id = ANY($1::uuid[]) AND brand_id = $2
        """

    # Tekrarları at ama sırayı koru — sayım karşılaştırması tekrarlı listede yanılmasın.
    requested = list(dict.fromkeys(document_ids))
    docs = await db.fetch(docs_sql, requested, brand_id)
    if len(docs) != len(requested):
        raise HTTPException(status_code=404, detail="Doküman bulunamadı")

    context_parts: list[str] = []
    remaining_chars = max_chars

    for doc in docs:
        doc_id = doc["id"]
        doc_name = doc["name"]

        if doc["raw_text"]:
            text = doc["raw_text"][:remaining_chars]
            context_parts.append(f"[{doc_name}]\n{text}")
            remaining_chars -= len(text)
        else:
            query_embedding = await _generate_embedding(prompt)
            if query_embedding is not None:
                chunks = await db.fetch(
                    f"""
                    SELECT content
                    FROM {chunks_table}
                    WHERE document_id = $1
                    ORDER BY embedding <=> $2::vector
                    LIMIT 5
                    """,
                    doc_id,
                    str(query_embedding),
                )
            else:
                chunks = await db.fetch(
                    f"""
                    SELECT content
                    FROM {chunks_table}
                    WHERE document_id = $1
                    ORDER BY chunk_index
                    LIMIT 5
                    """,
                    doc_id,
                )

            if chunks:
                combined = "\n\n".join(c["content"] for c in chunks)[:remaining_chars]
                context_parts.append(f"[{doc_name}]\n{combined}")
                remaining_chars -= len(combined)

        if remaining_chars <= 0:
            break

    return "\n\n---\n\n".join(context_parts)


async def get_product_document_context(
    product_ids: list[UUID],
    prompt: str,
    db: asyncpg.Connection,
    max_chars: int = 6_000,
    *,
    brand_id: UUID,
) -> str:
    """Ürün ID listesinden RAG context derle. Her ürüne bağlı tüm dokümanlar çekilir.

    Sprint 7 caption-gen ürün doküman bağlamı için.

    `brand_id` ZORUNLU: ürünün de dokümanlarının da bu markaya ait olması gerekir.
    Çağıranlardan biri (kısa video 2. aşama) ürün satırını sahiplik filtresiyle çekip
    sonucu KULLANMIYORDU — satır boş dönse bile doküman çağrısı koşuyordu, yani yabancı
    ürünün dokümanları yine enjekte ediliyordu (security review 2026-08-26, S2).
    Kapsamı çağırana bırakmak bu sınıfı kapatmıyor; kapı burada.
    """
    if not product_ids:
        return ""

    requested = list(dict.fromkeys(product_ids))
    owned = await db.fetch(
        "SELECT id FROM social.brand_products WHERE id = ANY($1::uuid[]) AND brand_id = $2",
        requested,
        brand_id,
    )
    if len(owned) != len(requested):
        raise HTTPException(status_code=404, detail="Ürün bulunamadı")

    rows = await db.fetch(
        "SELECT id FROM social.product_documents WHERE product_id = ANY($1::uuid[])",
        requested,
    )
    doc_ids = [r["id"] for r in rows]
    if not doc_ids:
        return ""
    return await get_document_context(
        doc_ids, prompt, db, max_chars=max_chars, kind="product", brand_id=brand_id
    )
